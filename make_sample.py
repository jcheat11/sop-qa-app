"""One-off script to generate a sample.docx for testing extract.py.
Not part of the app; safe to delete once you have a real SOP to test with.
"""

import docx
from docx.shared import Inches
import io
from PIL import Image


def _tiny_png_bytes() -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), color=(200, 0, 0)).save(buf, format="PNG")
    buf.seek(0)
    return buf


def main():
    doc = docx.Document()

    doc.add_heading("Standard Operating Procedure: Guest Checkout", level=1)

    doc.add_heading("Purpose", level=2)
    doc.add_paragraph(
        "This SOP describes the steps front-desk staff must follow when "
        "checking out a guest."
    )

    doc.add_heading("Step 1: Verify Folio", level=2)
    doc.add_paragraph("1. Open the property management system (PMS).")
    doc.add_paragraph("2. Search for the guest's reservation by last name or room number.")
    doc.add_paragraph("3. Confirm all charges on the folio are correct.")

    doc.add_heading("Step 2: Collect Payment", level=2)
    doc.add_paragraph("1. Ask the guest how they would like to settle the balance.")
    doc.add_paragraph("2. Process payment in the PMS.")
    p = doc.add_paragraph("3. A screenshot of the payment screen is shown below:")
    p.add_run().add_picture(_tiny_png_bytes(), width=Inches(1))

    doc.add_heading("Step 3: Room Status", level=2)
    doc.add_paragraph("1. Set the room status to 'Dirty' in the PMS.")
    doc.add_paragraph("2. Notify housekeeping via the daily board.")

    doc.add_heading("Escalation Table", level=2)
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Issue"
    table.rows[0].cells[1].text = "Contact"
    table.rows[1].cells[0].text = "Payment declined"
    table.rows[1].cells[1].text = "Front Desk Manager"
    table.rows[2].cells[0].text = "Guest dispute"
    table.rows[2].cells[1].text = "General Manager"

    doc.save("sample.docx")
    print("Wrote sample.docx")


if __name__ == "__main__":
    main()
